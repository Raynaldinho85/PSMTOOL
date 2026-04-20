from __future__ import annotations

import base64
from io import BytesIO

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import pytest
from pptx import Presentation

from psm_tool.core.metrics import compute_psm_kpis
from psm_tool.core.turnover_index import compute_profit_proxy, compute_turnover_index
from psm_tool.plots.render_static import (
    BrowserPreflightError,
    check_kaleido_browser,
    figure_to_png_bytes,
    prepare_figure_for_static_export,
)
from psm_tool.report.excel_export import build_excel_report
from psm_tool.report.pptx_builder import (
    TITLE_FONT_SIZE_PT,
    TITLE_MIN_FONT_SIZE_PT,
    _chart_label_overrides,
    _contained_rect,
    _fit_headline_for_pptx,
    _headline_from_sentence,
    _png_ratio,
    _truncate_text,
    build_pptx_report,
)

TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def _sample_curves() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "price": [10.0, 20.0, 30.0, 40.0],
            "too_cheap": [100.0, 80.0, 40.0, 10.0],
            "bargain": [100.0, 90.0, 55.0, 20.0],
            "expensive": [0.0, 20.0, 50.0, 80.0],
            "too_expensive": [0.0, 15.0, 45.0, 85.0],
            "not_bargain": [0.0, 10.0, 45.0, 80.0],
            "not_expensive": [100.0, 80.0, 50.0, 20.0],
        }
    )


def _sample_payload(*, with_cost: bool = False) -> dict:
    curves = _sample_curves()
    kpi_result = compute_psm_kpis(curves)
    turnover_result = compute_turnover_index(
        curves["price"].to_numpy(dtype=float),
        np.array([25.0, 60.0, 50.0, 20.0], dtype=float),
    )
    analysis = {
        "product_id": "Classic",
        "segment": "DE",
        "currency": "EUR",
        "curves": curves,
        "kpi_result": kpi_result,
        "kpis": kpi_result.as_dict(),
        "outlier_settings": {
            "enabled": True,
            "level": "medium",
            "q_low": 0.01,
            "q_high": 0.99,
        },
        "outlier_stats": {
            "excluded_n": 2,
            "analysis_n_after_outlier": 38,
        },
        "turnover_index_result": turnover_result,
    }
    if with_cost:
        profit_result = compute_profit_proxy(
            curves["price"].to_numpy(dtype=float),
            turnover_result.df["purchase_intention_pct"].to_numpy(dtype=float),
            unit_cost=12.0,
        )
        analysis["profit_proxy_result"] = profit_result
        analysis["unit_cost"] = 12.0
    return {"analyses": [analysis]}


def test_plotly_png_render_succeeds_when_browser_available() -> None:
    try:
        check_kaleido_browser()
    except BrowserPreflightError as exc:
        pytest.skip(f"Browser unavailable for kaleido: {exc}")

    fig = go.Figure(data=[go.Scatter(x=[0, 1], y=[0, 1])])
    payload = figure_to_png_bytes(fig)
    assert payload.startswith(b"\x89PNG")
    assert len(payload) > 100


def test_static_export_preparation_adds_safe_chart_margins() -> None:
    fig = go.Figure(data=[go.Scatter(x=[0, 1], y=[0, 1])])
    fig.update_layout(margin={"l": 5, "r": 10, "t": 20, "b": 15})

    prepared = prepare_figure_for_static_export(fig)

    assert prepared.layout.margin.l >= 70
    assert prepared.layout.margin.r >= 70
    assert prepared.layout.margin.t >= 120
    assert prepared.layout.margin.b >= 40
    assert fig.layout.margin.t == 20


def test_static_export_preparation_can_keep_exact_figure_margins() -> None:
    fig = go.Figure(data=[go.Scatter(x=[0, 1], y=[0, 1])])
    fig.update_layout(margin={"l": 0, "r": 0, "t": 0, "b": 0})

    prepared = prepare_figure_for_static_export(fig, safe_margins=False)

    assert prepared.layout.margin.l == 0
    assert prepared.layout.margin.r == 0
    assert prepared.layout.margin.t == 0
    assert prepared.layout.margin.b == 0


def test_pptx_contained_rect_preserves_image_ratio() -> None:
    x, y, width, height = _contained_rect(
        x=1.0,
        y=2.0,
        width=8.0,
        height=2.0,
        image_ratio=16.0 / 9.0,
    )

    assert x > 1.0
    assert y == 2.0
    assert height == 2.0
    assert round(width / height, 6) == round(16.0 / 9.0, 6)


def test_png_ratio_reads_png_header() -> None:
    png = (
        b"\x89PNG\r\n\x1a\n"
        + b"\x00\x00\x00\rIHDR"
        + (1600).to_bytes(4, "big")
        + (900).to_bytes(4, "big")
    )

    assert round(_png_ratio(png), 6) == round(16.0 / 9.0, 6)


def test_pptx_export_builds_report() -> None:
    try:
        check_kaleido_browser()
    except BrowserPreflightError as exc:
        pytest.skip(f"Browser unavailable for kaleido: {exc}")

    pptx_bytes = build_pptx_report(_sample_payload())
    assert pptx_bytes.startswith(b"PK")

    presentation = Presentation(BytesIO(pptx_bytes))
    assert len(presentation.slides) >= 1
    text_chunks = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_chunks.append(shape.text)
    full_text = "\n".join(text_chunks)
    assert "Model suggests" in full_text
    assert "Perception | Classic / DE (EUR)" in full_text
    assert "Economics proxy | Classic / DE (EUR)" in full_text


def test_pptx_chart_label_overrides_are_read_by_chart_id() -> None:
    analysis = {
        "marker_label_side_overrides": {
            "psm": {"pmi": "left", "opp": "auto"},
            "nms": {"max_revenue": "right"},
        }
    }

    assert _chart_label_overrides(analysis, "psm") == {"pmi": "left"}
    assert _chart_label_overrides(analysis, "nms") == {"max_revenue": "right"}
    assert _chart_label_overrides(analysis, "turnover") == {}


def test_pptx_chart_builders_receive_marker_overrides(monkeypatch) -> None:
    captured = {}

    def fake_psm_figure(*args, label_side_overrides=None, **kwargs):
        captured["psm"] = label_side_overrides
        return go.Figure()

    def fake_turnover_figure(*args, label_side_overrides=None, **kwargs):
        captured["turnover"] = label_side_overrides
        return go.Figure()

    monkeypatch.setattr("psm_tool.report.pptx_builder.make_psm_figure", fake_psm_figure)
    monkeypatch.setattr(
        "psm_tool.report.pptx_builder.make_turnover_index_figure",
        fake_turnover_figure,
    )
    monkeypatch.setattr("psm_tool.report.pptx_builder.figure_to_png_bytes", lambda fig: TINY_PNG)
    monkeypatch.setattr(
        "psm_tool.report.pptx_builder.kpi_summary_png_bytes",
        lambda analysis: TINY_PNG,
    )

    payload = _sample_payload()
    payload["analyses"][0]["marker_label_side_overrides"] = {
        "psm": {"pmi": "left"},
        "turnover": {"max_turnover_price": "right"},
    }

    build_pptx_report(payload)

    assert captured["psm"] == {"pmi": "left"}
    assert captured["turnover"] == {"max_turnover_price": "right"}


def test_pptx_headline_fitting_wraps_and_reduces_before_shortening() -> None:
    headline = (
        "This headline is intentionally long enough to need a slightly smaller font "
        "while still fitting as wrapped PowerPoint title text"
    )

    fitted, font_size = _fit_headline_for_pptx(headline)

    assert fitted == headline
    assert font_size < TITLE_FONT_SIZE_PT
    assert font_size >= TITLE_MIN_FONT_SIZE_PT
    assert not fitted.endswith("...")


def test_pptx_headline_shortening_uses_phrase_boundaries() -> None:
    headline = (
        "The acceptable range remains stable under current assumptions with additional "
        "diagnostic detail that would otherwise overrun the available PowerPoint title box"
    )

    fitted, font_size = _fit_headline_for_pptx(headline)

    assert fitted == "The acceptable range remains stable."
    assert font_size == TITLE_MIN_FONT_SIZE_PT
    assert not fitted.endswith("...")
    assert not fitted.endswith(("under...", "under curre...", "The acceptable r..."))


def test_pptx_headline_last_resort_shortening_does_not_cut_mid_word() -> None:
    headline = (
        "A deliberately long headline made of meaningful words that has no low information "
        "trailing connector and therefore must fall back to a clean whole word ending before "
        "the title box becomes too crowded"
    )

    fitted, _ = _fit_headline_for_pptx(headline)
    without_ellipsis = fitted.removesuffix("...")

    assert fitted.endswith("...")
    assert without_ellipsis
    assert headline.startswith(without_ellipsis)
    assert headline[len(without_ellipsis) :].startswith(" ")


def test_pptx_headline_source_no_longer_clips_before_presentation_fitting() -> None:
    headline = (
        "The acceptable range remains semantically complete even when the presentation "
        "layer later has to decide how to fit it into the title box"
    )

    assert _headline_from_sentence(headline, "Fallback") == headline
    assert _truncate_text("abcdefghijklmnopqrstuvwxyz", 10) == "abcdefg..."


def test_pptx_export_includes_profit_sentence_only_when_cost_is_provided() -> None:
    try:
        check_kaleido_browser()
    except BrowserPreflightError as exc:
        pytest.skip(f"Browser unavailable for kaleido: {exc}")

    with_cost = Presentation(BytesIO(build_pptx_report(_sample_payload(with_cost=True))))
    without_cost = Presentation(BytesIO(build_pptx_report(_sample_payload(with_cost=False))))

    with_text = "\n".join(
        shape.text for slide in with_cost.slides for shape in slide.shapes if hasattr(shape, "text")
    )
    without_text = "\n".join(
        shape.text
        for slide in without_cost.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    assert "Given unit cost" in with_text
    assert "highest profit proxy is achieved at" in with_text
    assert "Given unit cost" not in without_text


def test_pptx_export_adds_all_available_slides_for_each_analysis() -> None:
    try:
        check_kaleido_browser()
    except BrowserPreflightError as exc:
        pytest.skip(f"Browser unavailable for kaleido: {exc}")

    payload = _sample_payload(with_cost=True)
    analysis = payload["analyses"][0]
    payload["analyses"] = [analysis, {**analysis, "product_id": "Premium", "segment": "CH"}]

    presentation = Presentation(BytesIO(build_pptx_report(payload)))
    # Per analysis: KPI Summary + PSM + Turnover + Profit
    # (+ NMS only when available; sample payload has no NMS)
    assert len(presentation.slides) == 8


def test_pptx_export_shapes_stay_within_slide_bounds() -> None:
    try:
        check_kaleido_browser()
    except BrowserPreflightError as exc:
        pytest.skip(f"Browser unavailable for kaleido: {exc}")

    presentation = Presentation(BytesIO(build_pptx_report(_sample_payload(with_cost=True))))
    for slide in presentation.slides:
        slide_width = int(presentation.slide_width)
        slide_height = int(presentation.slide_height)
        for shape in slide.shapes:
            assert int(shape.left) >= 0
            assert int(shape.top) >= 0
            assert int(shape.left + shape.width) <= slide_width
            assert int(shape.top + shape.height) <= slide_height


def test_excel_export_contains_expected_sheets() -> None:
    excel_bytes = build_excel_report(_sample_payload(with_cost=True))
    workbook = pd.ExcelFile(BytesIO(excel_bytes))
    assert "kpis" in workbook.sheet_names
    assert "curves" in workbook.sheet_names
    assert "turnover_index" in workbook.sheet_names
    kpis = pd.read_excel(BytesIO(excel_bytes), sheet_name="kpis")
    assert {
        "outlier_filter_applied",
        "outlier_level",
        "outlier_excluded_n",
        "analysis_n_after_outlier",
    }.issubset(kpis.columns)
    nms_kpis = pd.read_excel(BytesIO(excel_bytes), sheet_name="nms_kpis")
    assert {"max_profit_price", "unit_cost"}.issubset(nms_kpis.columns)
    turnover = pd.read_excel(BytesIO(excel_bytes), sheet_name="turnover_index")
    assert {"profit_proxy_per_100", "profit_index"}.issubset(turnover.columns)
