from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt

from psm_tool.i18n.runtime import normalize_language, tr
from psm_tool.plots.benchmarks import PriceBenchmark, is_valid_tested_price
from psm_tool.plots.nms_plot import make_nms_figure
from psm_tool.plots.psm_plot import make_psm_figure
from psm_tool.plots.render_static import figure_to_png_bytes
from psm_tool.plots.turnover_index_plot import make_pi_economics_figure, make_turnover_index_figure
from psm_tool.report.insights import (
    build_nms_summary,
    build_profit_summary,
    build_psm_summary,
    build_turnover_summary,
)
from psm_tool.report.kpi_summary_png import kpi_summary_png_bytes
from psm_tool.report.wording_policy import apply_wording_policy, can_recommend

EMU_PER_INCH = 914400.0
CHART_EXPORT_RATIO = 16.0 / 9.0
TITLE_FONT_SIZE_PT = 23
TITLE_MIN_FONT_SIZE_PT = 21
TITLE_FIT_MIN_FONT_SIZE_PT = 17
CONTEXT_FONT_SIZE_PT = 11
SUMMARY_FONT_SIZE_PT = 11
SUMMARY_MIN_FONT_SIZE_PT = 9
SIDE_KPI_FONT_SIZE_PT = 11
SIDE_KPI_MIN_FONT_SIZE_PT = 9
TITLE_WRAP_CHAR_CAPACITY = 96
TITLE_REDUCED_WRAP_CHAR_CAPACITY = 122
TITLE_SHORTEN_CHAR_CAPACITY = 132
TITLE_MIN_SEMANTIC_CHARS = 34
BULLET_MAX_CHARS = 140
SIDE_KPI_MAX_CHARS = 92
PPTX_FONT_FAMILY_CANDIDATES = ("Calibri", "Arial", "Aptos", "DejaVu Sans", "Liberation Sans")
HEADLINE_TRAILING_BOUNDARIES = (
    " under ",
    " with ",
    " based on ",
    " because ",
    " due to ",
    " given ",
    " while ",
    " when ",
    " for ",
)


def _new_presentation(template_path: str | Path | None = None) -> Presentation:
    if template_path and Path(template_path).exists():
        return Presentation(str(template_path))
    return Presentation()


def _emu_to_in(value: int) -> float:
    return float(value) / EMU_PER_INCH


def _slide_size_in(presentation: Presentation) -> tuple[float, float]:
    return _emu_to_in(int(presentation.slide_width)), _emu_to_in(int(presentation.slide_height))


def _contained_rect(
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    image_ratio: float,
) -> tuple[float, float, float, float]:
    if width <= 0 or height <= 0:
        return x, y, max(width, 0.0), max(height, 0.0)
    target_ratio = width / height
    if target_ratio > image_ratio:
        fitted_h = height
        fitted_w = fitted_h * image_ratio
    else:
        fitted_w = width
        fitted_h = fitted_w / image_ratio
    return x + ((width - fitted_w) / 2.0), y + ((height - fitted_h) / 2.0), fitted_w, fitted_h


def _png_ratio(image_bytes: bytes, fallback: float = CHART_EXPORT_RATIO) -> float:
    if len(image_bytes) >= 24 and image_bytes[:8] == b"\x89PNG\r\n\x1a\n":
        width = int.from_bytes(image_bytes[16:20], "big")
        height = int.from_bytes(image_bytes[20:24], "big")
        if width > 0 and height > 0:
            return float(width) / float(height)
    return fallback


def _add_picture_contained(
    slide,
    image_bytes: bytes,
    *,
    x: float,
    y: float,
    width: float,
    height: float,
) -> None:
    image_x, image_y, image_w, image_h = _contained_rect(
        x=x,
        y=y,
        width=width,
        height=height,
        image_ratio=_png_ratio(image_bytes),
    )
    slide.shapes.add_picture(
        BytesIO(image_bytes),
        Inches(image_x),
        Inches(image_y),
        width=Inches(image_w),
        height=Inches(image_h),
    )


def _text_frame_defaults(frame) -> None:
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)


def _truncate_text(text: str, max_chars: int) -> str:
    clean = str(text).strip()
    if len(clean) <= max_chars:
        return clean
    return f"{clean[: max_chars - 3].rstrip(' ,;:.')}..."


def _clean_headline_text(text: str) -> str:
    return " ".join(str(text).strip().split())


def _as_complete_headline_phrase(text: str) -> str:
    clean = text.strip(" ,;:-")
    if clean and clean[-1] not in ".!?":
        return f"{clean}."
    return clean


def _phrase_boundary_headline(text: str, max_chars: int) -> str | None:
    lower_text = text.lower()
    for boundary in HEADLINE_TRAILING_BOUNDARIES:
        start = 0
        while True:
            idx = lower_text.find(boundary, start)
            if idx < 0:
                break
            candidate = text[:idx]
            if TITLE_MIN_SEMANTIC_CHARS <= len(candidate.strip()) <= max_chars:
                return _as_complete_headline_phrase(candidate)
            start = idx + len(boundary)

    candidates: list[str] = []
    for separator in (". ", "; ", " - ", " – ", " — "):
        start = 0
        while True:
            idx = text.find(separator, start)
            if idx < 0:
                break
            end = idx + 1 if separator in {". ", "; "} else idx
            candidates.append(text[:end])
            start = idx + len(separator)

    complete_candidates = [
        _as_complete_headline_phrase(candidate)
        for candidate in candidates
        if TITLE_MIN_SEMANTIC_CHARS <= len(candidate.strip()) <= max_chars
    ]
    if not complete_candidates:
        return None
    return max(complete_candidates, key=len)


def _word_boundary_headline(text: str, max_chars: int) -> str:
    limit = max_chars - 3
    split_at = text.rfind(" ", 0, limit + 1)
    if split_at < TITLE_MIN_SEMANTIC_CHARS:
        return text
    clean = text[:split_at].strip(" ,;:-")
    return f"{clean}..."


def _semantic_truncate_text(text: str, max_chars: int) -> str:
    clean = _clean_headline_text(text)
    if len(clean) <= max_chars:
        return clean
    return _phrase_boundary_headline(clean, max_chars) or _word_boundary_headline(clean, max_chars)


def _wrap_text_two_lines(text: str, *, max_chars_per_line: int) -> str | None:
    clean = _clean_headline_text(text)
    if not clean:
        return clean

    words = clean.split()
    if len(clean) <= max_chars_per_line:
        return clean

    candidates: list[tuple[int, str, str]] = []
    for split_idx in range(1, len(words)):
        first_line = " ".join(words[:split_idx]).strip()
        second_line = " ".join(words[split_idx:]).strip()
        if not first_line or not second_line:
            continue
        if len(first_line) > max_chars_per_line or len(second_line) > max_chars_per_line:
            continue
        candidates.append((max(len(first_line), len(second_line)), first_line, second_line))

    if not candidates:
        return None

    _line_score, first_line, second_line = min(candidates, key=lambda item: item[0])
    return f"{first_line}\n{second_line}"


def _fit_headline_for_pptx(text: str) -> tuple[str, int]:
    clean = _clean_headline_text(text)
    if len(clean) <= TITLE_WRAP_CHAR_CAPACITY:
        return clean, TITLE_FONT_SIZE_PT
    if len(clean) <= TITLE_REDUCED_WRAP_CHAR_CAPACITY:
        return clean, TITLE_FONT_SIZE_PT - 1
    if len(clean) <= TITLE_SHORTEN_CHAR_CAPACITY:
        return clean, TITLE_MIN_FONT_SIZE_PT
    return (
        _phrase_boundary_headline(clean, TITLE_SHORTEN_CHAR_CAPACITY)
        or _word_boundary_headline(clean, TITLE_SHORTEN_CHAR_CAPACITY),
        TITLE_MIN_FONT_SIZE_PT,
    )


def _configure_paragraph_layout(paragraph) -> None:
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0


def _set_text_frame_paragraphs(frame, paragraphs: list[str]) -> None:
    frame.clear()
    for idx, paragraph_text in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = paragraph_text
        _configure_paragraph_layout(paragraph)


def _best_fit_font_choice(frame, max_size: int) -> tuple[str, int] | None:
    for family in PPTX_FONT_FAMILY_CANDIDATES:
        try:
            size = frame._best_fit_font_size(family, max_size, False, False, None)
        except Exception:
            continue
        return family, size
    return None


def _apply_font_size(frame, family: str, size: int) -> None:
    try:
        frame._apply_fit(family, size, False, False)
    except Exception:
        frame.word_wrap = True
    for paragraph in frame.paragraphs:
        paragraph.font.name = family
        paragraph.font.size = Pt(size)
        for run in paragraph.runs:
            run.font.name = family
            run.font.size = Pt(size)


def _fit_text_frame_candidates(
    frame,
    *,
    paragraph_candidates: list[list[str]],
    max_size: int,
    min_size: int,
) -> int:
    if not paragraph_candidates:
        return max_size

    fallback_choice: tuple[list[str], tuple[str, int] | None] | None = None
    for paragraphs in paragraph_candidates:
        _set_text_frame_paragraphs(frame, paragraphs)
        choice = _best_fit_font_choice(frame, max_size)
        if choice is None:
            fallback_choice = (paragraphs, None)
            continue
        family, fitted_size = choice
        fallback_choice = (paragraphs, choice)
        if fitted_size >= min_size:
            _apply_font_size(frame, family, fitted_size)
            return fitted_size

    final_paragraphs, final_choice = fallback_choice or (paragraph_candidates[-1], None)
    _set_text_frame_paragraphs(frame, final_paragraphs)
    if final_choice is not None:
        family, fitted_size = final_choice
        _apply_font_size(frame, family, max(1, fitted_size))
        return fitted_size
    _apply_font_size(frame, PPTX_FONT_FAMILY_CANDIDATES[0], min_size)
    return min_size


def _add_title(slide, *, x: float, y: float, width: float, text: str) -> float:
    title_h = 0.84
    title_box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(title_h))
    frame = title_box.text_frame
    _text_frame_defaults(frame)
    full_text = _clean_headline_text(text)
    shortened_text, _font_size = _fit_headline_for_pptx(text)
    candidates = [[full_text]]
    if shortened_text != full_text:
        candidates.append([shortened_text])
    fitted_size = _fit_text_frame_candidates(
        frame,
        paragraph_candidates=candidates,
        max_size=TITLE_FONT_SIZE_PT,
        min_size=TITLE_FIT_MIN_FONT_SIZE_PT,
    )
    frame.paragraphs[0].font.size = Pt(fitted_size)
    frame.paragraphs[0].font.bold = True
    return title_h


def _add_context_line(
    slide,
    *,
    x: float,
    y: float,
    width: float,
    analysis: dict[str, Any],
    lens: str,
) -> float:
    language = normalize_language(analysis.get("language"))
    context_h = 0.34
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(context_h))
    frame = box.text_frame
    _text_frame_defaults(frame)
    frame.text = _truncate_text(
        tr(
            "{lens} | {product_id} / {segment} ({currency})",
            language,
            lens=lens,
            product_id=analysis["product_id"],
            segment=analysis["segment"],
            currency=analysis["currency"],
        ),
        120,
    )
    frame.paragraphs[0].font.size = Pt(CONTEXT_FONT_SIZE_PT)
    return context_h


def _add_bullet_text(
    slide,
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    sentences: list[str],
    max_items: int,
    font_size: int,
    strip_leading_label: str | None = None,
    preserve_font_size: bool = False,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    _text_frame_defaults(frame)
    label_prefix = f"{_clean_headline_text(strip_leading_label)}:" if strip_leading_label else None

    def _summary_line(sentence: str) -> str:
        clean = _clean_headline_text(sentence)
        if label_prefix and clean.startswith(label_prefix):
            clean = clean[len(label_prefix) :].strip()
        return clean

    summary_font_size = min(font_size, SUMMARY_FONT_SIZE_PT)
    summary_lines = [_summary_line(sentence) for sentence in sentences[:max_items]]
    bullet_lines = [f"- {line}" for line in summary_lines]
    wrapped_two_line_candidates = [
        [
            (
                f"- {wrapped_line}"
                if (
                    wrapped_line := _wrap_text_two_lines(
                        line,
                        max_chars_per_line=max_chars,
                    )
                )
                is not None
                else f"- {line}"
            )
            for line in summary_lines
        ]
        for max_chars in (78, 72, 66, 60)
    ]
    fallback_lines = [
        f"- {_semantic_truncate_text(line, BULLET_MAX_CHARS)}" for line in summary_lines
    ]
    tighter_lines = [f"- {_semantic_truncate_text(line, 110)}" for line in summary_lines]
    _fit_text_frame_candidates(
        frame,
        paragraph_candidates=[
            bullet_lines,
            *wrapped_two_line_candidates,
            fallback_lines,
            tighter_lines,
        ],
        max_size=summary_font_size,
        min_size=summary_font_size if preserve_font_size else SUMMARY_MIN_FONT_SIZE_PT,
    )


def _headline_from_sentence(sentence: str, fallback: str) -> str:
    text = _clean_headline_text(sentence)
    if not text:
        return fallback
    return text


def _build_turnover_summary_safe(
    *,
    turnover_result: Any,
    currency: str,
    segment_label: str,
    source: str | None,
    kpi_statuses: dict[str, str] | None = None,
    language: str | None = None,
) -> list[str]:
    try:
        return build_turnover_summary(
            turnover_result,
            currency=currency,
            segment_label=segment_label,
            source=source,
            kpi_statuses=kpi_statuses,
            language=language,
        )
    except TypeError as exc:
        if "unexpected keyword argument 'kpi_statuses'" not in str(exc):
            raise
        return build_turnover_summary(
            turnover_result,
            currency=currency,
            segment_label=segment_label,
            source=source,
            language=language,
        )


def _build_profit_summary_safe(
    *,
    profit_result: Any,
    currency: str,
    segment_label: str,
    product_label: str,
    kpi_statuses: dict[str, str] | None = None,
    language: str | None = None,
) -> list[str]:
    try:
        return build_profit_summary(
            profit_result,
            currency=currency,
            segment_label=segment_label,
            product_label=product_label,
            kpi_statuses=kpi_statuses,
            language=language,
        )
    except TypeError as exc:
        if "unexpected keyword argument 'kpi_statuses'" not in str(exc):
            raise
        return build_profit_summary(
            profit_result,
            currency=currency,
            segment_label=segment_label,
            product_label=product_label,
            language=language,
        )


def _status_map(analysis: dict[str, Any], keys: tuple[str, ...]) -> dict[str, str]:
    kpis = analysis.get("kpis", {})
    return {key: str(kpis.get(f"{key}_status", "closest")) for key in keys}


def _tested_price_benchmarks(analysis: dict[str, Any]) -> list[PriceBenchmark]:
    if not bool(analysis.get("tested_price_active", False)):
        return []
    tested_price = analysis.get("tested_price")
    if not is_valid_tested_price(tested_price):
        return []
    language = normalize_language(analysis.get("language"))
    return [
        PriceBenchmark(
            label=tr("Tested Price", language),
            price=float(tested_price),
            key="tested_price",
        )
    ]


def _chart_label_overrides(analysis: dict[str, Any], chart_id: str) -> dict[str, str]:
    overrides = analysis.get("marker_label_side_overrides", {})
    if not isinstance(overrides, dict):
        return {}
    chart_overrides = overrides.get(chart_id, {})
    if not isinstance(chart_overrides, dict):
        return {}
    return {
        str(marker_key): str(side)
        for marker_key, side in chart_overrides.items()
        if str(side) in {"left", "right"}
    }


def _psm_action_title(analysis: dict[str, Any]) -> str:
    language = normalize_language(analysis.get("language"))
    kpis = analysis["kpis"]
    statuses = _status_map(analysis, ("pmi", "opp", "idp", "pme"))
    recommendation_allowed = can_recommend(statuses, allow_interval=False)
    if recommendation_allowed:
        sentence = apply_wording_policy(
            (
                f"Preis im akzeptierten Preisbereich {kpis['accepted_low']:.2f}-"
                f"{kpis['accepted_high']:.2f} {analysis['currency']} ansetzen."
                if language == "de"
                else (
                    f"Set price in accepted range {kpis['accepted_low']:.2f}-"
                    f"{kpis['accepted_high']:.2f} {analysis['currency']}."
                )
            ),
            lens=tr("Perception", language),
            language=language,
        )
    else:
        sentence = apply_wording_policy(
            tr("OPP recommendation is blocked because PSM intersections are not clean.", language),
            lens=tr("Perception", language),
            status_flags={"unstable": True, "recommendation_blocked": True},
            language=language,
        )
    return _headline_from_sentence(
        sentence,
        tr("Perception: Model suggests pricing guidance.", language),
    )


def _turnover_action_title(analysis: dict[str, Any]) -> str:
    language = normalize_language(analysis.get("language"))
    turnover = analysis.get("turnover_index_result")
    if turnover is None:
        return tr("Economics proxy: Model suggests turnover diagnostics only.", language)
    statuses = _status_map(analysis, ("opp",))
    recommendation_allowed = can_recommend(statuses, allow_interval=False)
    if recommendation_allowed:
        sentence = apply_wording_policy(
            (
                f"Turnover nahe {turnover.max_turnover_price:.2f} "
                f"{analysis['currency']} priorisieren."
                if language == "de"
                else (
                    f"Maximize turnover near "
                    f"{turnover.max_turnover_price:.2f} {analysis['currency']}."
                )
            ),
            lens=tr("Economics proxy", language),
            language=language,
        )
    else:
        sentence = apply_wording_policy(
            tr(
                "Turnover target-price recommendation is blocked due to non-clean intersections.",
                language,
            ),
            lens=tr("Economics proxy", language),
            status_flags={"unstable": True, "recommendation_blocked": True},
            language=language,
        )
    return _headline_from_sentence(
        sentence,
        tr("Economics proxy: Model suggests turnover guidance.", language),
    )


def _nms_action_title(analysis: dict[str, Any]) -> str:
    language = normalize_language(analysis.get("language"))
    nms_result = analysis.get("nms_result")
    if nms_result is None:
        return tr("Modeled demand: Model suggests NMS diagnostics.", language)
    sentence = apply_wording_policy(
        (
            "Trial "
            f"({nms_result.max_trial_price:.2f}) und Revenue ({nms_result.max_revenue_price:.2f}) "
            f"in {analysis['currency']} ausbalancieren."
            if language == "de"
            else (
                f"Balance trial ({nms_result.max_trial_price:.2f}) and revenue "
                f"({nms_result.max_revenue_price:.2f}) in {analysis['currency']}."
            )
        ),
        lens=tr("Modeled demand", language),
        language=language,
    )
    return _headline_from_sentence(
        sentence,
        tr("Modeled demand: Model suggests trial/revenue context.", language),
    )


def _profit_action_title(analysis: dict[str, Any]) -> str:
    language = normalize_language(analysis.get("language"))
    profit = analysis.get("profit_proxy_result")
    if profit is None:
        return tr("Economics proxy: Model suggests profit diagnostics only.", language)
    statuses = _status_map(analysis, ("pmi", "pme"))
    recommendation_allowed = can_recommend(statuses, allow_interval=False)
    if recommendation_allowed:
        sentence = apply_wording_policy(
            (
                f"Profit-Proxy nahe {profit.max_profit_price:.2f} "
                f"{analysis['currency']} optimieren."
                if language == "de"
                else (
                    f"Optimize profit proxy near "
                    f"{profit.max_profit_price:.2f} {analysis['currency']}."
                )
            ),
            lens=tr("Economics proxy", language),
            language=language,
        )
    else:
        sentence = apply_wording_policy(
            tr(
                "Profit target-price recommendation is blocked due to non-clean intersections.",
                language,
            ),
            lens=tr("Economics proxy", language),
            status_flags={"unstable": True, "recommendation_blocked": True},
            language=language,
        )
    return _headline_from_sentence(
        sentence,
        tr("Economics proxy: Model suggests profit guidance.", language),
    )


def _add_kpi_summary(
    slide,
    *,
    x: float,
    y: float,
    width: float,
    height: float,
    analysis: dict[str, Any],
) -> None:
    language = normalize_language(analysis.get("language"))
    kpis = analysis["kpis"]
    currency = analysis["currency"]

    def _format_point(key: str) -> str:
        status = str(kpis.get(f"{key}_status", "closest"))
        value = float(kpis.get(key, 0.0))
        if status == "clean":
            return f"{currency} {value:.2f}"
        if status == "interval":
            low = kpis.get(f"{key}_low")
            high = kpis.get(f"{key}_high")
            if low is None or high is None:
                return f"{currency} {value:.2f} (~mid)"
            return (
                f"[{currency} {float(low):.2f}, {currency} {float(high):.2f}] "
                f"(~ {currency} {value:.2f})"
            )
        return f"{currency} {value:.2f} (diagnostic)"

    outlier_settings = analysis.get("outlier_settings", {})
    outlier_stats = analysis.get("outlier_stats", {})
    outlier_enabled = bool(outlier_settings.get("enabled", False))
    if outlier_enabled:
        level = str(outlier_settings.get("level", "medium")).capitalize()
        q_low = outlier_settings.get("q_low")
        q_high = outlier_settings.get("q_high")
        excluded = int(outlier_stats.get("excluded_n", 0))
        outlier_line = (
            tr(
                "Outlier filter: {level} ({q_low:.1f}%-{q_high:.1f}%), excluded n={excluded}",
                language,
                level=level,
                q_low=float(q_low) * 100.0,
                q_high=float(q_high) * 100.0,
                excluded=excluded,
            )
            if q_low is not None and q_high is not None
            else tr(
                "Outlier filter: {level}, excluded n={excluded}",
                language,
                level=level,
                excluded=excluded,
            )
        )
    else:
        outlier_line = tr("Outlier filter: Off", language)

    pmi_clean = str(kpis.get("pmi_status", "closest")) == "clean"
    pme_clean = str(kpis.get("pme_status", "closest")) == "clean"
    opp_clean = str(kpis.get("opp_status", "closest")) == "clean"
    idp_clean = str(kpis.get("idp_status", "closest")) == "clean"
    accepted_range_text = (
        f"{currency} {kpis['accepted_low']:.2f} - {currency} {kpis['accepted_high']:.2f}"
        if pmi_clean and pme_clean
        else "— (diagnostic)"
    )
    stress_text = (
        f"{kpis['price_stress']:.2f} [{kpis['stress_flag']}]"
        if opp_clean and idp_clean
        else "— (diagnostic)"
    )
    lines = [
        f"PMI: {_format_point('pmi')}",
        f"OPP: {_format_point('opp')}",
        f"IDP: {_format_point('idp')}",
        f"PME: {_format_point('pme')}",
        tr(
            "Accepted range: {accepted_range_text}",
            language,
            accepted_range_text=accepted_range_text,
        ),
        tr("Price stress (OPP-IDP): {stress_text}", language, stress_text=stress_text),
        outlier_line,
    ]
    if any(not is_clean for is_clean in (pmi_clean, opp_clean, idp_clean, pme_clean)):
        lines.append(tr("Intersection quality: interpret with caution.", language))

    text_box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = text_box.text_frame
    _text_frame_defaults(frame)
    full_lines = [_clean_headline_text(line) for line in lines]
    shortened_lines = [_semantic_truncate_text(line, SIDE_KPI_MAX_CHARS) for line in lines]
    _fit_text_frame_candidates(
        frame,
        paragraph_candidates=[full_lines, shortened_lines],
        max_size=SIDE_KPI_FONT_SIZE_PT,
        min_size=SIDE_KPI_MIN_FONT_SIZE_PT,
    )


def _add_kpi_summary_slide(presentation: Presentation, analysis: dict[str, Any]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_w, slide_h = _slide_size_in(presentation)
    margin = 0.35
    _add_picture_contained(
        slide,
        kpi_summary_png_bytes(analysis),
        x=margin,
        y=margin,
        width=slide_w - (2 * margin),
        height=slide_h - (2 * margin),
    )


def _add_psm_slide(presentation: Presentation, analysis: dict[str, Any]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_w, slide_h = _slide_size_in(presentation)
    margin = 0.42
    content_w = slide_w - (2 * margin)
    gap = 0.14

    title_h = _add_title(
        slide,
        x=margin,
        y=0.16,
        width=content_w,
        text=_psm_action_title(analysis),
    )
    context_y = 0.16 + title_h + 0.05
    context_h = _add_context_line(
        slide,
        x=margin,
        y=context_y,
        width=content_w,
        analysis=analysis,
        lens=tr("Perception", normalize_language(analysis.get("language"))),
    )

    summary_h = min(1.60, max(1.20, slide_h * 0.20))
    summary_y = slide_h - margin - summary_h
    chart_top = context_y + context_h + 0.07
    chart_h = max(2.45, summary_y - chart_top - 0.08)

    kpi_w = min(3.7, max(2.8, content_w * 0.29))
    chart_w = max(4.6, content_w - kpi_w - gap)
    if (chart_w + kpi_w + gap) > content_w:
        kpi_w = max(2.4, content_w - chart_w - gap)
    chart_x = margin
    kpi_x = chart_x + chart_w + gap

    figure = make_psm_figure(
        analysis["curves"],
        analysis["kpi_result"],
        price_benchmarks=_tested_price_benchmarks(analysis),
        label_side_overrides=_chart_label_overrides(analysis, "psm"),
        language=normalize_language(analysis.get("language")),
    )
    image_bytes = figure_to_png_bytes(figure)
    _add_picture_contained(
        slide,
        image_bytes,
        x=chart_x,
        y=chart_top,
        width=chart_w,
        height=chart_h,
    )

    _add_kpi_summary(
        slide,
        x=kpi_x,
        y=chart_top,
        width=max(2.2, content_w - chart_w - gap),
        height=chart_h,
        analysis=analysis,
    )
    _add_psm_summary_bullets(
        slide,
        analysis=analysis,
        x=margin,
        y=summary_y,
        width=content_w,
        height=summary_h,
    )


def _add_psm_summary_bullets(
    slide,
    *,
    analysis: dict[str, Any],
    x: float,
    y: float,
    width: float,
    height: float,
) -> None:
    language = normalize_language(analysis.get("language"))
    nms_result = analysis.get("nms_result")
    nms_kpis = None
    if nms_result is not None:
        nms_kpis = {
            "max_trial_price": nms_result.max_trial_price,
            "max_revenue_price": nms_result.max_revenue_price,
        }
    sentences = build_psm_summary(
        analysis["kpis"],
        currency=analysis["currency"],
        segment_label=str(analysis["segment"]),
        product_label=str(analysis["product_id"]),
        nms_kpis=nms_kpis,
        language=language,
    )
    _add_bullet_text(
        slide,
        x=x,
        y=y,
        width=width,
        height=height,
        sentences=sentences,
        max_items=4,
        font_size=12,
        strip_leading_label=tr("Perception", language),
        preserve_font_size=True,
    )


def _add_turnover_index_slide(presentation: Presentation, analysis: dict[str, Any]) -> bool:
    turnover_result = analysis.get("turnover_index_result")
    if turnover_result is None:
        return False

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_w, slide_h = _slide_size_in(presentation)
    margin = 0.42
    content_w = slide_w - (2 * margin)

    title_h = _add_title(
        slide,
        x=margin,
        y=0.16,
        width=content_w,
        text=_turnover_action_title(analysis),
    )
    context_y = 0.16 + title_h + 0.05
    context_h = _add_context_line(
        slide,
        x=margin,
        y=context_y,
        width=content_w,
        analysis=analysis,
        lens=tr("Economics proxy", normalize_language(analysis.get("language"))),
    )

    summary_h = min(1.50, max(1.10, slide_h * 0.19))
    summary_y = slide_h - margin - summary_h
    chart_top = context_y + context_h + 0.07
    chart_h = max(2.40, summary_y - chart_top - 0.08)

    figure = make_turnover_index_figure(
        turnover_result,
        currency=analysis["currency"],
        price_benchmarks=_tested_price_benchmarks(analysis),
        label_side_overrides=_chart_label_overrides(analysis, "turnover"),
        language=normalize_language(analysis.get("language")),
    )
    image_bytes = figure_to_png_bytes(figure)
    _add_picture_contained(
        slide,
        image_bytes,
        x=margin,
        y=chart_top,
        width=content_w,
        height=chart_h,
    )

    summary_sentences = _build_turnover_summary_safe(
        turnover_result=turnover_result,
        currency=analysis["currency"],
        segment_label=str(analysis["segment"]),
        source=analysis.get("turnover_source"),
        kpi_statuses=_status_map(analysis, ("opp",)),
        language=normalize_language(analysis.get("language")),
    )
    _add_bullet_text(
        slide,
        x=margin,
        y=summary_y,
        width=content_w,
        height=summary_h,
        sentences=summary_sentences,
        max_items=3,
        font_size=12,
        strip_leading_label=tr("Economics proxy", normalize_language(analysis.get("language"))),
        preserve_font_size=True,
    )
    return True


def _add_nms_slide(presentation: Presentation, analysis: dict[str, Any]) -> None:
    nms_result = analysis.get("nms_result")
    if nms_result is None:
        return

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_w, slide_h = _slide_size_in(presentation)
    margin = 0.42
    content_w = slide_w - (2 * margin)
    gap = 0.14

    title_h = _add_title(
        slide,
        x=margin,
        y=0.16,
        width=content_w,
        text=_nms_action_title(analysis),
    )
    context_y = 0.16 + title_h + 0.05
    context_h = _add_context_line(
        slide,
        x=margin,
        y=context_y,
        width=content_w,
        analysis=analysis,
        lens=tr("Modeled demand", normalize_language(analysis.get("language"))),
    )

    summary_h = min(1.50, max(1.10, slide_h * 0.19))
    summary_y = slide_h - margin - summary_h
    chart_top = context_y + context_h + 0.07
    chart_h = max(2.40, summary_y - chart_top - 0.08)

    meta_w = min(3.2, max(2.6, content_w * 0.27))
    chart_w = max(4.5, content_w - meta_w - gap)
    if (chart_w + meta_w + gap) > content_w:
        meta_w = max(2.2, content_w - chart_w - gap)

    figure = make_nms_figure(
        nms_result,
        price_benchmarks=_tested_price_benchmarks(analysis),
        label_side_overrides=_chart_label_overrides(analysis, "nms"),
        language=normalize_language(analysis.get("language")),
    )
    image_bytes = figure_to_png_bytes(figure)
    _add_picture_contained(
        slide,
        image_bytes,
        x=margin,
        y=chart_top,
        width=chart_w,
        height=chart_h,
    )

    box_x = margin + chart_w + gap
    box = slide.shapes.add_textbox(
        Inches(box_x),
        Inches(chart_top),
        Inches(meta_w),
        Inches(chart_h),
    )
    frame = box.text_frame
    _text_frame_defaults(frame)
    language = normalize_language(analysis.get("language"))
    lines = [
        tr(
            "Max Trial: {currency} {price:.2f}",
            language,
            currency=analysis["currency"],
            price=float(nms_result.max_trial_price),
        ),
        tr(
            "Max Revenue: {currency} {price:.2f}",
            language,
            currency=analysis["currency"],
            price=float(nms_result.max_revenue_price),
        ),
        tr("Included N: {included_n}", language, included_n=int(nms_result.included_n)),
    ]
    full_lines = [_clean_headline_text(line) for line in lines]
    shortened_lines = [_semantic_truncate_text(line, SIDE_KPI_MAX_CHARS) for line in lines]
    _fit_text_frame_candidates(
        frame,
        paragraph_candidates=[full_lines, shortened_lines],
        max_size=SIDE_KPI_FONT_SIZE_PT,
        min_size=SIDE_KPI_MIN_FONT_SIZE_PT,
    )

    _add_nms_summary_bullets(
        slide,
        analysis=analysis,
        nms_result=nms_result,
        x=margin,
        y=summary_y,
        width=content_w,
        height=summary_h,
    )


def _add_profit_slide(presentation: Presentation, analysis: dict[str, Any]) -> bool:
    turnover_result = analysis.get("turnover_index_result")
    profit_result = analysis.get("profit_proxy_result")
    unit_cost = analysis.get("unit_cost")
    if turnover_result is None or profit_result is None or unit_cost is None:
        return False

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_w, slide_h = _slide_size_in(presentation)
    margin = 0.42
    content_w = slide_w - (2 * margin)

    title_h = _add_title(
        slide,
        x=margin,
        y=0.16,
        width=content_w,
        text=_profit_action_title(analysis),
    )
    context_y = 0.16 + title_h + 0.05
    context_h = _add_context_line(
        slide,
        x=margin,
        y=context_y,
        width=content_w,
        analysis=analysis,
        lens=tr("Economics proxy", normalize_language(analysis.get("language"))),
    )

    summary_h = min(1.50, max(1.10, slide_h * 0.19))
    summary_y = slide_h - margin - summary_h
    chart_top = context_y + context_h + 0.07
    chart_h = max(2.40, summary_y - chart_top - 0.08)

    figure = make_pi_economics_figure(
        turnover_result,
        currency=analysis["currency"],
        mode="profit",
        profit_result=profit_result,
        unit_cost=float(unit_cost),
        price_benchmarks=_tested_price_benchmarks(analysis),
        language=normalize_language(analysis.get("language")),
    )
    image_bytes = figure_to_png_bytes(figure)
    _add_picture_contained(
        slide,
        image_bytes,
        x=margin,
        y=chart_top,
        width=content_w,
        height=chart_h,
    )

    summary = _build_profit_summary_safe(
        profit_result=profit_result,
        currency=analysis["currency"],
        segment_label=str(analysis["segment"]),
        product_label=str(analysis["product_id"]),
        kpi_statuses=_status_map(analysis, ("pmi", "pme")),
        language=normalize_language(analysis.get("language")),
    )
    _add_bullet_text(
        slide,
        x=margin,
        y=summary_y,
        width=content_w,
        height=summary_h,
        sentences=summary,
        max_items=3,
        font_size=12,
    )
    return True


def _add_nms_summary_bullets(
    slide,
    *,
    analysis: dict[str, Any],
    nms_result: Any,
    x: float,
    y: float,
    width: float,
    height: float,
) -> None:
    sentences = build_nms_summary(
        nms_result=nms_result,
        currency=analysis["currency"],
        segment_label=str(analysis["segment"]),
        language=normalize_language(analysis.get("language")),
    )
    _add_bullet_text(
        slide,
        x=x,
        y=y,
        width=width,
        height=height,
        sentences=sentences,
        max_items=4,
        font_size=12,
        strip_leading_label=tr("Modeled demand", normalize_language(analysis.get("language"))),
        preserve_font_size=True,
    )


def build_pptx_report(
    report_payload: dict[str, Any],
    template_path: str | Path | None = None,
) -> bytes:
    presentation = _new_presentation(template_path=template_path)
    analyses = report_payload.get("analyses", [])

    for analysis in analyses:
        _add_kpi_summary_slide(presentation, analysis)
        _add_psm_slide(presentation, analysis)
        _add_turnover_index_slide(presentation, analysis)
        _add_profit_slide(presentation, analysis)
        _add_nms_slide(presentation, analysis)

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
